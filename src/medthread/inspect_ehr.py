
from pathlib import Path
from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[2]

EXCEL_PATH = (
    PROJECT_ROOT
    / "data"
    / "sample_records"
    / "synthetic_ehr.xlsx"
)


def inspect_workbook():
    workbook = load_workbook(
        EXCEL_PATH,
        read_only=True,
        data_only=True,
    )

    try:
        print("\nMEDTHREAD AI — EHR DATA INSPECTION")
        print("-" * 45)

        print("Available sheets:", workbook.sheetnames)

        worksheet = workbook["Encounters"]

        rows = worksheet.iter_rows(values_only=True)
        headers = next(rows)

        print("\nNumber of columns:", len(headers))
        print("Column names:")

        for column in headers:
            print(" -", column)

        print("\nSample encounters:")

        patient_index = headers.index("Patient ID")
        encounter_index = headers.index("Encounter ID")

        for row in rows:
            patient_id = row[patient_index]
            encounter_id = row[encounter_index]
            if patient_id is None or encounter_id is None:
                continue

            patient_id = f"{int(patient_id):06d}"

            print(
                f"Patient: {patient_id} | "
                f"Encounter: {encounter_id}"
            )

    finally:
        workbook.close()


if __name__ == "__main__":
    inspect_workbook()
