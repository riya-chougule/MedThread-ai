
from io import BytesIO
import re

from openpyxl import load_workbook


def get_patient_ids(file_bytes: bytes) -> list[str]:
    """Return unique patient IDs from all workbook sheets."""
    workbook = load_workbook(
        filename=BytesIO(file_bytes),
        read_only=True,
        data_only=True,
    )

    patient_ids = set()

    try:
        for worksheet in workbook.worksheets:
            rows = worksheet.iter_rows(values_only=True)
            headers = next(rows, None)

            if not headers:
                continue

            header_names = [
                str(value).strip().lower() if value is not None else ""
                for value in headers
            ]

            patient_id_index = next(
                (
                    i for i, name in enumerate(header_names)
                    if name in {"patient id", "patient identifier"}
                ),
                None,
            )

            if patient_id_index is None:
                continue

            for row in rows:
                if patient_id_index < len(row):
                    value = row[patient_id_index]
                    if value is not None and str(value).strip():
                        patient_ids.add(str(value).strip())

    finally:
        workbook.close()

    return sorted(patient_ids)


def excel_to_record_text(
    file_bytes: bytes,
    selected_patient_id: str,
) -> str:
    """Convert only the selected patient's rows into traceable text."""
    workbook = load_workbook(
        filename=BytesIO(file_bytes),
        read_only=True,
        data_only=True,
    )

    documents = []

    try:
        for worksheet in workbook.worksheets:
            rows = worksheet.iter_rows(values_only=True)
            headers = next(rows, None)

            if not headers:
                continue

            clean_headers = [
                str(value).strip() if value is not None else f"Column {i + 1}"
                for i, value in enumerate(headers)
            ]

            patient_id_index = next(
                (
                    i for i, header in enumerate(clean_headers)
                    if header.lower() in {"patient id", "patient identifier"}
                ),
                None,
            )

            if patient_id_index is None:
                continue

            sheet_slug = re.sub(
                r"[^A-Z0-9]+",
                "-",
                worksheet.title.upper(),
            ).strip("-") or "SHEET"

            for row_number, row in enumerate(rows, start=2):
                if patient_id_index >= len(row):
                    continue

                row_patient_id = row[patient_id_index]
                if row_patient_id is None:
                    continue

                if str(row_patient_id).strip() != selected_patient_id:
                    continue

                row_sections = []

                for header, value in zip(clean_headers, row):
                    if value is None:
                        continue

                    if hasattr(value, "isoformat"):
                        value = value.isoformat()

                    section = re.sub(r"[^A-Za-z -]", " ", header)
                    section = re.sub(r"\s+", " ", section).strip()

                    if not section or not section[0].isalpha():
                        section = f"Field {header}"

                    row_sections.append(f"{section}:\n{value}")

                if not row_sections:
                    continue

                document_id = (
                    f"EXCEL-{sheet_slug}-ROW-{row_number}"
                )

                documents.append(
                    f"## Document {len(documents) + 1}: "
                    f"Excel sheet {worksheet.title}, row {row_number}\n\n"
                    f"Document ID: {document_id}\n\n"
                    + "\n\n".join(row_sections)
                )

    finally:
        workbook.close()

    if not documents:
        raise ValueError(
            f"No records found for patient {selected_patient_id}."
        )

    return "\n\n".join(documents)
